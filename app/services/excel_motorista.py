"""Geração do modelo de Excel e leitura/validação da planilha preenchida
para o cadastro de Motorista (REST, ver docs/API_MAPEAMENTO.md).

Bloco básico espelha a aba "Relação de motoristas" da planilha oficial de
coleta da Apisul (mesmo arquivo usado pro Ponto Geográfico); bloco avançado
expõe campos que a API aceita mas não fazem parte da coleta padrão.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

NOME_ABA_DADOS = "Motorista"
NOME_ABA_INSTRUCOES = "Instrucoes"
MARCADOR_LINHA_EXEMPLO = "EXEMPLO-APAGAR-ESTA-LINHA"

NOME_ABA_COLETA_OFICIAL = "Relação de motoristas"
LINHA_CABECALHO_COLETA_OFICIAL = 6
MARCADOR_LINHA_EXEMPLO_COLETA_OFICIAL = "exemplo"

# Confirmado pela Apisul em 2026-09-18.
TIPOS_MOTORISTA: dict[str, int] = {
    "Autônomo": 1,
    "Agregado": 2,
    "Frota": 3,
}

# Confirmado pela Apisul em 2026-09-18. Documento de identidade por país.
TIPOS_DOCUMENTO: dict[str, int] = {
    "CPF": 1,
    "DNI": 2,
    "RUT": 3,
    "CI": 4,
}

# PENDENTE: swagger só diz que é um inteiro (enum 1-7), não achamos a tabela
# oficial ainda. Todo dado real que vimos usa "Celular" — assumindo código 1
# até a Apisul confirmar (ver docs/API_MAPEAMENTO.md).
TIPOS_TELEFONE: dict[str, int] = {
    "Celular": 1,
}


@dataclass
class Coluna:
    campo: str
    titulo: str
    obrigatorio: bool
    tipo: str  # "texto" | "inteiro" | "decimal" | "data" | "tipo_motorista" | "tipo_documento" | "tipo_telefone"
    exemplo: Any
    descricao: str
    avancado: bool = False


COLUNAS: list[Coluna] = [
    Coluna("NomeCompleto", "Nome", True, "texto", "João da Silva", ""),
    Coluna("Documento", "CPF", True, "texto", "62530423601", "Só números."),
    Coluna("TipoMotorista", "Tipo Motorista", True, "tipo_motorista", "Frota",
           "Selecione: Autônomo, Agregado ou Frota."),
    Coluna("Cidade", "Cidade", True, "texto", "Porto Alegre", ""),
    Coluna("UF", "Estado", True, "texto", "RS", "Sigla com 2 letras."),
    Coluna("Pais", "País", False, "texto", "Brasil", "Padrão: Brasil, se deixado em branco."),
    Coluna("TipoTelefone", "Tipo Telefone", False, "tipo_telefone", "Celular", ""),
    Coluna("NumeroTelefone", "Número Telefone", True, "texto", "(51) 99999-0099",
           "Com DDD — o app separa DDD do restante do número."),
    # --- bloco avançado ---
    Coluna("TipoDocumento", "[Avançado] Tipo Documento", False, "tipo_documento", "",
           "CPF (padrão), DNI (Argentina), RUT (Chile) ou CI (Bolívia/Paraguai/Uruguai).", avancado=True),
    Coluna("PlacaVeiculoPadrao", "[Avançado] Placa Veículo Padrão", False, "texto", "", "", avancado=True),
    Coluna("CPFGestorFrota", "[Avançado] CPF Gestor Frota", False, "texto", "", "", avancado=True),
    Coluna("LocalizadorApisulMob", "[Avançado] Localizador ApisulMob", False, "texto", "", "", avancado=True),
    Coluna("DataAdmissao", "[Avançado] Data Admissão", False, "data", "", "", avancado=True),
    Coluna("CodContratante", "[Avançado] Cód. Contratante", False, "texto", "", "", avancado=True),
    Coluna("CodFuncionario", "[Avançado] Cód. Funcionário", False, "texto", "", "", avancado=True),
]

COLUNAS_POR_CAMPO = {c.campo: c for c in COLUNAS}

CAMPOS_COLETA_OFICIAL = [
    "NomeCompleto", "Documento", "TipoMotorista", "Cidade", "UF", "Pais", "TipoTelefone", "NumeroTelefone",
]


def gerar_modelo() -> BytesIO:
    wb = Workbook()
    aba: Worksheet = wb.active
    aba.title = NOME_ABA_DADOS

    cabecalho_font = Font(bold=True, color="FFFFFF")
    cabecalho_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    avancado_fill = PatternFill(start_color="5B7A99", end_color="5B7A99", fill_type="solid")
    obrigatorio_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for col_idx, coluna in enumerate(COLUNAS, start=1):
        celula = aba.cell(row=1, column=col_idx, value=coluna.titulo + (" *" if coluna.obrigatorio else ""))
        celula.font = cabecalho_font
        celula.fill = avancado_fill if coluna.avancado else cabecalho_fill
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if coluna.obrigatorio:
            aba.cell(row=2, column=col_idx).fill = obrigatorio_fill
        aba.column_dimensions[celula.column_letter].width = max(16, len(coluna.titulo) + 2)

    aba.cell(row=2, column=1, value=MARCADOR_LINHA_EXEMPLO)
    for col_idx, coluna in enumerate(COLUNAS, start=1):
        if col_idx == 1:
            continue
        aba.cell(row=2, column=col_idx, value=coluna.exemplo)
    for cel in aba[2]:
        cel.font = Font(italic=True, color="808080")

    aba.freeze_panes = "A3"

    for campo, opcoes in (("TipoMotorista", TIPOS_MOTORISTA), ("TipoTelefone", TIPOS_TELEFONE), ("TipoDocumento", TIPOS_DOCUMENTO)):
        idx = next(i for i, c in enumerate(COLUNAS, start=1) if c.campo == campo)
        letra = aba.cell(row=1, column=idx).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes.keys()) + '"', allow_blank=True, showDropDown=False)
        aba.add_data_validation(dv)
        dv.add(f"{letra}2:{letra}1048576")

    instrucoes = wb.create_sheet(NOME_ABA_INSTRUCOES)
    instrucoes.append(["Campo", "Obrigatório", "Tipo", "Exemplo", "Descrição"])
    for cel in instrucoes[1]:
        cel.font = cabecalho_font
        cel.fill = cabecalho_fill
    for coluna in COLUNAS:
        instrucoes.append([coluna.titulo, "Sim" if coluna.obrigatorio else "Não", coluna.tipo, coluna.exemplo, coluna.descricao])
    for largura, col_letra in zip([32, 14, 16, 20, 60], ["A", "B", "C", "D", "E"]):
        instrucoes.column_dimensions[col_letra].width = largura

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@dataclass
class LinhaValidada:
    numero_linha: int
    dados: dict[str, Any] = field(default_factory=dict)
    erros: list[str] = field(default_factory=list)

    @property
    def valida(self) -> bool:
        return not self.erros


def _converter(valor, coluna: Coluna, erros: list[str]):
    vazio = valor is None or (isinstance(valor, str) and not valor.strip())

    domínios = {"tipo_motorista": TIPOS_MOTORISTA, "tipo_documento": TIPOS_DOCUMENTO, "tipo_telefone": TIPOS_TELEFONE}
    if coluna.tipo in domínios:
        if vazio:
            if coluna.obrigatorio:
                erros.append(f"'{coluna.titulo}' é obrigatório.")
            return None
        texto = str(valor).strip()
        opcoes = domínios[coluna.tipo]
        if texto not in opcoes:
            erros.append(f"'{coluna.titulo}': valor '{texto}' não reconhecido. Use um de: {', '.join(opcoes)}.")
            return None
        return texto

    if vazio:
        if coluna.obrigatorio:
            erros.append(f"'{coluna.titulo}' é obrigatório.")
        return None
    try:
        if coluna.tipo == "texto":
            return str(valor).strip()
        if coluna.tipo == "inteiro":
            return int(float(valor))
        if coluna.tipo == "decimal":
            return float(valor)
        if coluna.tipo == "data":
            if isinstance(valor, datetime):
                return valor
            texto = str(valor).strip()
            for formato in ("%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
                try:
                    return datetime.strptime(texto, formato)
                except ValueError:
                    continue
            erros.append(f"'{coluna.titulo}': data '{valor}' em formato não reconhecido.")
            return None
    except (TypeError, ValueError):
        erros.append(f"'{coluna.titulo}': valor '{valor}' inválido para tipo {coluna.tipo}.")
        return None
    return valor


def _separar_ddd_numero(telefone: str) -> tuple[str | None, str | None]:
    digitos = re.sub(r"\D", "", telefone or "")
    if len(digitos) < 3:
        return None, digitos or None
    return digitos[:2], digitos[2:]


def _ler_aba_modelo_app(aba) -> list[LinhaValidada]:
    linhas: list[LinhaValidada] = []
    for numero_linha, linha_valores in enumerate(aba.iter_rows(min_row=2, values_only=True), start=2):
        if not any(linha_valores):
            continue
        if linha_valores and str(linha_valores[0] or "").strip() == MARCADOR_LINHA_EXEMPLO:
            continue

        lv = LinhaValidada(numero_linha=numero_linha)
        for idx, coluna in enumerate(COLUNAS):
            valor_bruto = linha_valores[idx] if idx < len(linha_valores) else None
            lv.dados[coluna.campo] = _converter(valor_bruto, coluna, lv.erros)
        linhas.append(lv)
    return linhas


def _ler_aba_coleta_oficial(aba) -> list[LinhaValidada]:
    linhas: list[LinhaValidada] = []
    for numero_linha, linha_valores in enumerate(
        aba.iter_rows(min_row=LINHA_CABECALHO_COLETA_OFICIAL + 1, values_only=True),
        start=LINHA_CABECALHO_COLETA_OFICIAL + 1,
    ):
        valores_sem_indice = linha_valores[1:] if linha_valores else ()
        if not any(valores_sem_indice):
            continue
        if str(linha_valores[0] or "").strip().lower() == MARCADOR_LINHA_EXEMPLO_COLETA_OFICIAL:
            continue

        lv = LinhaValidada(numero_linha=numero_linha)
        for idx, campo in enumerate(CAMPOS_COLETA_OFICIAL):
            coluna = COLUNAS_POR_CAMPO[campo]
            valor_bruto = valores_sem_indice[idx] if idx < len(valores_sem_indice) else None
            lv.dados[campo] = _converter(valor_bruto, coluna, lv.erros)
        linhas.append(lv)
    return linhas


def ler_planilha(arquivo) -> list[LinhaValidada]:
    """Lê um arquivo .xlsx (file-like). Aceita o modelo do app (aba
    'Motorista') ou a planilha oficial de coleta da Apisul (aba
    'Relação de motoristas')."""
    wb = load_workbook(arquivo, data_only=True)
    if NOME_ABA_DADOS in wb.sheetnames:
        return _ler_aba_modelo_app(wb[NOME_ABA_DADOS])
    if NOME_ABA_COLETA_OFICIAL in wb.sheetnames:
        return _ler_aba_coleta_oficial(wb[NOME_ABA_COLETA_OFICIAL])
    raise ValueError(
        f"A planilha enviada não tem nem a aba '{NOME_ABA_DADOS}' (modelo do app) "
        f"nem '{NOME_ABA_COLETA_OFICIAL}' (planilha oficial de coleta RQ IMP 002)."
    )


def _somente_digitos(valor: Any) -> Any:
    if not isinstance(valor, str):
        return valor
    digitos = "".join(c for c in valor if c.isdigit())
    return digitos or None


def montar_payload_rest(dados: dict[str, Any]) -> dict[str, Any]:
    """Converte uma linha validada no formato esperado por
    MotoristaModeloIntegracao (separa DDD do número, mapeia os rótulos de
    domínio para os códigos numéricos confirmados)."""
    ddd, numero = _separar_ddd_numero(dados.get("NumeroTelefone", ""))

    tipo_motorista = dados.get("TipoMotorista")
    tipo_documento = dados.get("TipoDocumento") or "CPF"
    tipo_telefone = dados.get("TipoTelefone") or "Celular"

    documento = dados.get("Documento")
    if tipo_documento == "CPF":
        # RUT/DNI/CI podem ter letra verificadora (ex.: RUT chileno) -- só
        # sanitiza pra so-digitos quando é CPF, que a API rejeita com erro
        # 1065 se vier pontuado (mesmo comportamento visto no CNPJ do
        # Ponto Geografico).
        documento = _somente_digitos(documento)

    payload = {
        "NomeCompleto": dados.get("NomeCompleto"),
        "Documento": documento,
        "TipoDocumento": TIPOS_DOCUMENTO.get(tipo_documento),
        "TipoMotorista": TIPOS_MOTORISTA.get(tipo_motorista),
        "DDD": ddd,
        "NumeroTelefonePrincipal": numero,
        "TipoTelefone": TIPOS_TELEFONE.get(tipo_telefone),
        "Ativo": True,
        "UF": dados.get("UF"),
        "Cidade": dados.get("Cidade"),
        "PlacaVeiculoPadrao": dados.get("PlacaVeiculoPadrao"),
        "CPFGestorFrota": dados.get("CPFGestorFrota"),
        "LocalizadorApisulMob": dados.get("LocalizadorApisulMob"),
        "DataAdmissao": dados.get("DataAdmissao").isoformat() if dados.get("DataAdmissao") else None,
        "CodContratante": dados.get("CodContratante"),
        "CodFuncionario": dados.get("CodFuncionario"),
    }
    return {k: v for k, v in payload.items() if v is not None}
